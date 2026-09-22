"""Le e reescreve SO os textos de Highlights de um .pptx ja gerado -
graficos, tabelas, logo e rodape ficam intactos.

Fluxo pensado pra analise: `sequencia.py` monta o deck com os
highlights automaticos (os mesmos da tela), voce abre no PowerPoint,
decide o que cada slide deveria dizer, e aplica os textos novos aqui
sem precisar regerar nada.

    uv run python highlights.py analise.pptx                  # mostra o que ha hoje
    uv run python highlights.py analise.pptx -a novos.txt     # aplica
    uv run python highlights.py analise.pptx -a novos.txt -o v2.pptx

Formato de `novos.txt` - um cabecalho por slide, so os slides que voce
quer mudar (os demais ficam como estao):

    ## 1
    N+NE responde por 18% do Valor com Presente do pais, mas cresce
    o dobro da media nacional.

    ## 3
    Outro texto...

Linhas em branco separam paragrafos dentro do mesmo slide.
"""

from __future__ import annotations

import argparse
import copy
import re
import sys

from pptx import Presentation
from pptx.util import Pt

# nome dado a caixa por `export_pptx._add_highlight_textbox`. Decks
# gerados antes disso nao tem nome nenhum, dai o fallback por conteudo.
_SHAPE_NAME = "symrise-highlight"
_HEADING = "Highlights"

# o texto automatico e limitado a 500 chars (insights.MAX_CHARS) e a
# caixa foi dimensionada pra isso - acima disso avisa, porque o
# PowerPoint nao reflui: o texto so vaza pra fora do slide
_SOFT_LIMIT = 500


def _highlight_box(slide):
    """A caixa de Highlights do slide, ou None se ele nao tiver uma
    (slides de continuacao de tabela, por exemplo)."""
    for shape in slide.shapes:
        if shape.name == _SHAPE_NAME:
            return shape
    # fallback pra decks gerados antes da shape ganhar nome
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.paragraphs:
            first = shape.text_frame.paragraphs[0].text.strip()
            if first == _HEADING:
                return shape
    return None


def _body_paragraphs(tf):
    """Os paragrafos de corpo (tudo menos o cabecalho "Highlights")."""
    return tf.paragraphs[1:]


def read_all(prs) -> list[tuple[int, str | None]]:
    out = []
    for n, slide in enumerate(prs.slides, start=1):
        box = _highlight_box(slide)
        if box is None:
            out.append((n, None))
        else:
            texto = "\n".join(p.text for p in _body_paragraphs(box.text_frame))
            out.append((n, texto.strip()))
    return out


def _set_text(box, novo: str) -> None:
    """Troca o corpo preservando a formatacao. Atribuir
    `text_frame.text` colapsaria tudo num unico run sem estilo (fonte,
    corpo e cor voltariam pro default do PowerPoint), entao o texto vai
    run a run, e paragrafos extras sao clonados do primeiro - herdando
    fonte/tamanho/cor/espacamento ja calibrados."""
    tf = box.text_frame
    corpo = _body_paragraphs(tf)
    if not corpo:
        raise RuntimeError("caixa de Highlights sem paragrafo de corpo")

    linhas = [ln.strip() for ln in novo.strip().split("\n\n")]
    linhas = [ln for ln in linhas if ln] or [""]

    modelo = corpo[0]

    # sobras de uma aplicacao anterior com mais paragrafos que esta
    for p in corpo[len(linhas):]:
        p._p.getparent().remove(p._p)

    # paragrafos que faltam: clone XML do primeiro (leva pPr + rPr)
    for _ in range(len(linhas) - len(corpo)):
        novo_p = copy.deepcopy(modelo._p)
        modelo._p.getparent().append(novo_p)

    for p, texto in zip(_body_paragraphs(tf), linhas):
        runs = p.runs
        if not runs:  # paragrafo vazio herdado: cria um run com o estilo do modelo
            p._p.append(copy.deepcopy(modelo.runs[0]._r))
            runs = p.runs
        runs[0].text = texto
        for extra in runs[1:]:  # o texto novo cabe todo no primeiro run
            extra._r.getparent().remove(extra._r)


def parse_updates(texto: str) -> dict[int, str]:
    """Le o formato `## <numero do slide>` seguido do texto."""
    updates: dict[int, str] = {}
    atual: int | None = None
    buf: list[str] = []

    for linha in texto.splitlines():
        m = re.match(r"^##\s*(?:slide\s*)?(\d+)\s*$", linha.strip(), re.IGNORECASE)
        if m:
            if atual is not None:
                updates[atual] = "\n".join(buf).strip()
            atual = int(m.group(1))
            if atual in updates:
                raise SystemExit(f"slide {atual} aparece duas vezes no arquivo de textos")
            buf = []
        elif atual is not None:
            buf.append(linha)
        elif linha.strip():
            raise SystemExit(f"texto fora de qualquer slide (falta um '## N' antes): {linha.strip()[:60]!r}")

    if atual is not None:
        updates[atual] = "\n".join(buf).strip()
    if not updates:
        raise SystemExit("nenhum bloco '## N' encontrado no arquivo de textos")
    return updates


def apply(path_in: str, path_updates: str, path_out: str) -> None:
    prs = Presentation(path_in)
    with open(path_updates, encoding="utf-8") as fh:
        updates = parse_updates(fh.read())

    slides = list(prs.slides)
    for n in sorted(updates):
        if not 1 <= n <= len(slides):
            raise SystemExit(f"slide {n} nao existe - o deck tem {len(slides)}")
        if _highlight_box(slides[n - 1]) is None:
            raise SystemExit(f"slide {n} nao tem caixa de Highlights (slide de continuacao de tabela?)")

    for n in sorted(updates):
        novo = updates[n]
        _set_text(_highlight_box(slides[n - 1]), novo)
        aviso = f"  [!] {len(novo)} chars, acima de {_SOFT_LIMIT} - confira se nao vazou da caixa" if len(novo) > _SOFT_LIMIT else ""
        print(f"slide {n}: {len(novo)} chars{aviso}")

    prs.save(path_out)
    print(f"\n{len(updates)} slide(s) atualizado(s) -> {path_out}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", help="deck gerado por sequencia.py")
    ap.add_argument("-a", "--apply", metavar="TXT", help="arquivo com os textos novos")
    ap.add_argument("-o", "--out", help="saida (padrao: sobrescreve o proprio deck)")
    args = ap.parse_args()

    if not args.apply:
        for n, texto in read_all(Presentation(args.pptx)):
            if texto is None:
                print(f"## {n}\n(sem caixa de Highlights)\n")
            else:
                print(f"## {n}\n{texto}\n")
        return

    apply(args.pptx, args.apply, args.out or args.pptx)


if __name__ == "__main__":
    main()
